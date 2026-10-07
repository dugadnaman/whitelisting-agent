"""Consumer-visible campaign preview behavior: identity, mapping, and no writes."""

import csv
import io
import json
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from api import app
from moengage_preview import prepare_batch, rows_from_file, rows_from_jira_brief

USER = {"sub": "user-1", "tenant_id": "tata", "role": "operator", "email": "owner@example.com"}
CATALOG = {
    "account": "tata", "workspace_id": "workspace-tata",
    "segments": [{"id": "seg-vip", "name": "VIP"}, {"id": "seg-new", "name": "New customers"}],
    "subscription_categories": ["Offers"],
    "email_senders": [{"from_address": "mail@example.com", "sender_name": "Tata",
                       "connector_type": "SENDGRID", "connector_name": "Primary"}],
    "push_platforms": [
        {"platform": "ANDROID", "notification_channel": "promotions"},
        {"platform": "IOS", "ios_flags": {"send_to_all_eligible_device": True,
                                         "exclude_provisional_push_devices": False,
                                         "send_to_only_provisional_push_enabled_devices": False}},
        {"platform": "WEB"},
    ],
    "assets": [{"id": "asset-1", "url": "https://example.com/banner.png"}],
    "email_templates": [{"id": "tmpl-1"}],
    "whatsapp_senders": [{"sender_name": "Tata Capital Financial Services Limited", "phone_number": "+919876543210", "provider": "KARIX"}],
    "whatsapp_templates": [{"id": "festive_wa_1", "name": "festive_wa_1"}],
}
BASE = {
    "account": "tata", "source_ref": "brief.json", "row_id": "1", "channel": "EMAIL",
    "campaign_name": "October VIP", "segment_id": "seg-vip",
    "scheduled_at": "2026-10-15T10:00:00+05:30", "timezone": "Asia/Kolkata",
    "content_type": "PROMOTIONAL", "subscription_category": "Offers",
    "from_address": "mail@example.com", "subject": "Exclusive offer",
    "html_content": "<p>Welcome VIP</p>",
}


def preview(rows, catalog=CATALOG):
    return prepare_batch(rows, "tata", catalog, USER["email"], source_type="json")


def test_mixed_rows_preserve_identity_and_produce_exact_candidate():
    push = {
        "account": "tata", "source_ref": "brief.json", "row_id": "2", "channel": "PUSH",
        "campaign_name": "App offer", "segment_name": "New customers",
        "scheduled_at": BASE["scheduled_at"], "timezone": BASE["timezone"],
        "push_platform": "ANDROID", "push_title": "A new offer", "push_message": "See details",
        "click_url": "https://example.com/offer", "push_image_asset_id": "asset-1",
    }
    other = {**BASE, "row_id": "3", "account": "bajaj"}
    unknown = {**BASE, "row_id": "4", "segment_id": "seg-missing"}
    result = preview([BASE, push, other, unknown])
    assert (result["ready"], result["blocked"], result["write_eligible"]) == (2, 2, False)
    assert [(item["source_ref"], item["row_id"]) for item in result["items"]] == [
        ("brief.json", "1"), ("brief.json", "2"), ("brief.json", "3"), ("brief.json", "4")
    ]
    email_payload = result["items"][0]["candidate_v5_payload"]
    assert email_payload == {
        "channel": "EMAIL", "campaign_delivery_type": "ONE_TIME", "created_by": USER["email"],
        "basic_details": {"name": "October VIP", "content_type": "PROMOTIONAL", "subscription_category": "Offers"},
        "connector": {"connector_type": "SENDGRID", "connector_name": "Primary"},
        "campaign_content": {"content": {"email": {
            "html_content": "<p>Welcome VIP</p>", "subject": "Exclusive offer",
            "sender_name": "Tata", "from_address": "mail@example.com",
            "reply_to_address": "mail@example.com",
        }}},
        "segmentation_details": {"included_filters": {"filter_operator": "and", "filters": [
            {"filter_type": "custom_segments", "name": "VIP", "id": "seg-vip"}
        ]}},
        "scheduling_details": {"delivery_type": "AT_FIXED_TIME", "start_time": "2026-10-15T10:00:00", "timezone": "Asia/Kolkata"},
    }
    push_payload = result["items"][1]["candidate_v5_payload"]
    assert push_payload["campaign_content"]["content"]["push"]["android"] == {
        "template_type": "BASIC", "basic_details": {
            "title": "A new offer", "message": "See details", "notification_channel": "promotions",
            "default_click_action": "DEEPLINKING", "default_click_action_value": "https://example.com/offer",
            "image_url": "https://example.com/banner.png",
        }
    }
    assert result["items"][2]["candidate_v5_payload"] is None
    assert "Row account/workspace does not match" in result["items"][2]["issues"][0]
    assert result["items"][3]["status"] == "blocked"


def test_ambiguous_audience_duplicate_ids_and_schedule_fail_closed():
    catalog = {**CATALOG, "segments": CATALOG["segments"] + [{"id": "seg-other", "name": "VIP"}]}
    rows = [
        {**BASE, "row_id": "1", "segment_id": "", "segment_name": "VIP"},
        {**BASE, "row_id": "1"},
        {**BASE, "row_id": "3", "segment_id": "", "segment_name": ""},
        {**BASE, "row_id": "4", "scheduled_at": "2026-10-15T10:00:00", "timezone": "Asia/Kolkata"},
        {**BASE, "row_id": "5", "scheduled_at": "2026-10-15T10:00:00+00:00", "timezone": "Asia/Kolkata"},
        {**BASE, "row_id": "6", "created_by": "spoof@example.com"},
    ]
    result = preview(rows, catalog)
    assert result["ready"] == 0
    assert all(item["candidate_v5_payload"] is None for item in result["items"])
    assert "Duplicate source_ref/row_id in batch" in result["items"][0]["issues"]
    assert "Duplicate source_ref/row_id in batch" in result["items"][1]["issues"]
    assert "never all users" in result["items"][2]["issues"][0]
    assert "ISO offset" in result["items"][3]["issues"][0]
    assert "ISO offset" in result["items"][4]["issues"][0]
    assert "creator does not match" in result["items"][5]["issues"][0]


def test_catalog_assets_and_platforms_require_exact_reviewed_values():
    for platform in ("IOS", "WEB"):
        row = {**{key: value for key, value in BASE.items() if key not in (
            "content_type", "subscription_category", "from_address", "subject", "html_content",
        )}, "channel": "PUSH", "push_platform": platform, "push_title": "Hi",
               "push_message": "Check this", "click_url": "https://example.com", "row_id": platform}
        item = preview([row])["items"][0]
        assert item["status"] == "preview_ready"
        assert platform.lower() in item["candidate_v5_payload"]["campaign_content"]["content"]["push"]
    template = {**BASE, "html_content": "", "email_template_id": "tmpl-1", "subject": ""}
    assert preview([template])["items"][0]["candidate_v5_payload"]["campaign_content"]["content"]["email"]["custom_template_id"] == "tmpl-1"
    invalid = {**BASE, "email_attachment_ids": ["unknown-attachment"]}
    assert preview([invalid])["items"][0]["status"] == "blocked"
    approved = {**BASE, "email_attachment_ids": ["asset-1"]}
    attachments = preview([approved])["items"][0]["candidate_v5_payload"]["campaign_content"]["content"]["email"]["attachments"]
    assert attachments == [{"file_type": "URL", "url": "https://example.com/banner.png"}]
    mixed = {**BASE, "push_message": "Ignore me"}
    assert "another channel" in preview([mixed])["items"][0]["issues"][0]


def test_csv_and_xlsx_retaining_worksheet_and_multiline_line_numbers():
    buffer = io.StringIO(newline="")
    no_id = {key: value for key, value in BASE.items() if key not in ("source_ref", "row_id")}
    writer = csv.DictWriter(buffer, fieldnames=list(no_id))
    writer.writeheader()
    writer.writerow({**no_id, "html_content": "<p>First\nline</p>"})
    writer.writerow(no_id)
    csv_rows = rows_from_file("../batch.csv", buffer.getvalue().encode(), "tata")
    assert [row["row_id"] for row in csv_rows] == ["3", "4"]
    assert [row["source_ref"] for row in csv_rows] == ["batch.csv", "batch.csv"]
    workbook = Workbook()
    first = workbook.active
    first.title = "Email"
    first.append(list(no_id))
    first.append(list(no_id.values()))
    second = workbook.create_sheet("Push")
    second.append(list(no_id))
    second.append(list(no_id.values()))
    data = io.BytesIO()
    workbook.save(data)
    xlsx_rows = rows_from_file("batch.xlsx", data.getvalue(), "tata")
    assert [row["row_id"] for row in xlsx_rows] == ["Email!2", "Push!2"]


def test_parsed_jira_attachments_keep_provenance_without_invented_copy():
    brief = {
        "issue_key": "SWCM-200", "summary": "Lending offer", "account": "tata",
        "email_templates": [{"body": "Hello & goodbye", "subject": "Hello"}],
        "channel_counts": {"email": 1, "push": 1},
        "moengage_campaign": {"push_title": "Guessed", "push_body": "Guessed", "status": "DRAFT"},
        "attachments_mapped": [{"id": "a1", "filename": "Mailer.zip", "local_path": "/private/file"}],
    }
    overrides = {"campaign_name": "Reviewed offer", "segment_id": "seg-vip",
                 "scheduled_at": BASE["scheduled_at"], "timezone": BASE["timezone"],
                 "content_type": "PROMOTIONAL", "subscription_category": "Offers",
                 "from_address": "mail@example.com", "push_platform": "ANDROID"}
    rows = rows_from_jira_brief(brief, overrides, "tata")
    result = preview(rows)
    assert [item["row_id"] for item in result["items"]] == ["email:1", "push:1"]
    assert result["ready"] == 1 and result["blocked"] == 1
    assert result["items"][0]["candidate_v5_payload"]["campaign_content"]["content"]["email"]["html_content"] == "<p>Hello &amp; goodbye</p>"
    assert result["items"][0]["source_attachments"] == [{"id": "a1", "filename": "Mailer.zip"}]
    assert result["items"][1]["candidate_v5_payload"] is None
    assert brief["moengage_campaign"]["push_body"] not in json.dumps(result)

def test_jira_brief_whatsapp_templates_generate_v5_candidate_drafts_and_tata_subaccounts():
    brief = {
        "issue_key": "SWCM-106",
        "summary": "WhatsApp festival brief",
        "account": "tcl_promo",
        "whatsapp_templates": [
            {"template_name": "festive_wa_1", "body": "Dear {{1}}, festival offer"}
        ],
        "channel_counts": {"whatsapp": 1},
        "attachments_mapped": [{"id": "att-1", "filename": "copy.xlsx"}],
    }
    overrides = {
        "segment_id": "seg-vip",
        "scheduled_at": BASE["scheduled_at"],
        "timezone": BASE["timezone"],
    }
    # Allows 'tcl_promo' brief to generate drafts under parent 'tata' workspace
    rows = rows_from_jira_brief(brief, overrides, "tata")
    assert len(rows) == 1
    assert rows[0]["row_id"] == "whatsapp:1"
    assert rows[0]["channel"] == "WHATSAPP"
    assert rows[0]["whatsapp_template_id"] == "festive_wa_1"
    assert rows[0]["campaign_name"] == "festive_wa_1"

    batch = preview(rows)
    assert batch["ready"] == 1
    assert batch["blocked"] == 0
    candidate = batch["items"][0]["candidate_v5_payload"]
    assert candidate["channel"] == "WHATSAPP"
    assert candidate["basic_details"]["name"] == "festive_wa_1"
    assert candidate["campaign_content"]["content"]["whatsapp"]["template_id"] == "festive_wa_1"
    assert candidate["campaign_content"]["content"]["whatsapp"]["sender_name"] == "Tata Capital Financial Services Limited"

def test_empty_batch_bad_file_and_wrong_workspace_are_rejected():
    with pytest.raises(ValueError, match="At least one"):
        preview([])
    with pytest.raises(ValueError, match="workspace catalog"):
        preview([BASE], {**CATALOG, "account": "bajaj"})
    with pytest.raises(ValueError, match="Invalid XLSX"):
        rows_from_file("bad.xlsx", b"not an XLSX file", "tata")
    with pytest.raises(ValueError, match="unique"):
        rows_from_file("bad.csv", b"account,account\nbajaj,tata\n", "tata")


def test_preview_routes_are_authenticated_and_never_call_providers(provision_user):
    _, headers = provision_user(tenant="tata", email=USER["email"])
    client = TestClient(app)
    with patch("requests.request", side_effect=AssertionError("no provider calls")):
        denied = client.post("/api/moengage/drafts/preview", json={"account": "tata", "catalog": CATALOG, "row": BASE})
        assert denied.status_code == 401
        denied = client.post("/api/moengage/drafts/preview", headers=headers, json={"account": "bajaj", "catalog": CATALOG, "row": BASE})
        assert denied.status_code == 403
        allowed = client.post("/api/moengage/drafts/preview", headers=headers, json={"account": "tata", "catalog": CATALOG, "row": BASE})
        assert allowed.status_code == 200
        assert allowed.json()["items"][0]["status"] == "preview_ready"
        assert allowed.json()["write_eligible"] is False
        data = io.StringIO()
        csv.writer(data).writerows([list({k: v for k, v in BASE.items() if k not in ("row_id", "source_ref")}),
                                    list(v for k, v in BASE.items() if k not in ("row_id", "source_ref"))])
        file_result = client.post(
            "/api/moengage/drafts/preview-file",
            headers=headers,
            data={"account": "tata", "catalog_json": json.dumps(CATALOG)},
            files={"file": ("campaigns.csv", data.getvalue().encode(), "text/csv")},
        )
        assert file_result.status_code == 200
        assert file_result.json()["items"][0]["row_id"] == "2"
        assert file_result.json()["ready"] == 1
        brief = {"issue_key": "SWCM-200", "account": "tata", "channel_counts": {"push": 1}}
        jira = client.post("/api/moengage/drafts/preview", headers=headers, json={"account": "tata", "catalog": CATALOG, "brief": brief})
        assert jira.status_code == 200
        assert jira.json()["items"][0]["status"] == "blocked"


def test_parse_jira_execution_datetime():
    from moengage_preview import parse_jira_execution_datetime

    assert parse_jira_execution_datetime("25.09.2026 & 11:00 AM") == "2026-09-25T11:00:00+05:30"
    assert parse_jira_execution_datetime("25-09-2026 11:00 AM") == "2026-09-25T11:00:00+05:30"
    assert parse_jira_execution_datetime("25/09/2026 2:30 PM") == "2026-09-25T14:30:00+05:30"
    assert parse_jira_execution_datetime("01-10-2026 @ 09:00") == "2026-10-01T09:00:00+05:30"
    assert parse_jira_execution_datetime("2026-10-06") == "2026-10-06T11:00:00+05:30"
    assert parse_jira_execution_datetime("immediate") is None
    assert parse_jira_execution_datetime("asap") is None
    assert parse_jira_execution_datetime("N/A") is None


def test_jira_brief_with_kv_table_extracts_campaign_fields():
    brief = {
        "issue_key": "SWCM-105",
        "summary": "TCLService_CSFD_EWP Revamp - Communication Campaign - Corporate Services portal - Revamp",
        "account": "tcl_promo",
        "email_templates": [{"html_content": "<p>New Portal is Live</p>", "template_name": "Portal_Mailer"}],
        "channel_counts": {"email": 1},
        "campaign_metadata": {
            "campaign_name": "CSFD_EWP_Corporate_portal_Sept’26",
            "scheduled_at": "2026-09-25T11:00:00+05:30",
            "timezone": "Asia/Kolkata",
            "email_subject": "Our New Portal is Live!",
        },
    }
    overrides = {
        "segment_id": "seg-vip",
        "from_address": "mail@example.com",
    }
    rows = rows_from_jira_brief(brief, overrides, "tata")
    assert len(rows) == 1
    row = rows[0]
    assert row["campaign_name"] == "CSFD_EWP_Corporate_portal_Sept’26"
    assert row["scheduled_at"] == "2026-09-25T11:00:00+05:30"
    assert row["timezone"] == "Asia/Kolkata"
    assert row["subject"] == "Our New Portal is Live!"
    assert row["content_type"] == "PROMOTIONAL"
    assert row["subscription_category"] == "Offers"

    batch = preview(rows)
    assert batch["ready"] == 1
    assert batch["blocked"] == 0
    candidate = batch["items"][0]["candidate_v5_payload"]
    assert candidate["basic_details"]["name"] == "CSFD_EWP_Corporate_portal_Sept’26"
    assert candidate["scheduling_details"]["start_time"] == "2026-09-25T11:00:00"
    assert candidate["scheduling_details"]["timezone"] == "Asia/Kolkata"
    assert candidate["campaign_content"]["content"]["email"]["subject"] == "Our New Portal is Live!"


def test_jira_brief_without_segment_override_blocks_only_audience():
    brief = {
        "issue_key": "SWCM-105",
        "summary": "TCLService_CSFD_EWP Revamp - Communication Campaign - Corporate Services portal - Revamp",
        "account": "tcl_promo",
        "description_text": (
            "Campaign Name | CSFD_EWP_Corporate_portal_Sept’26\n"
            "Date & Time of execution | 25.09.2026 & 11:00 AM\n"
            "Email Text | Subject: Our New Portal is Live!\n"
        ),
        "email_templates": [{"html_content": "<p>New Portal is Live</p>", "template_name": "Portal_Mailer"}],
        "channel_counts": {"email": 1},
    }
    # Operator provides from_address but forgets segment_id
    overrides = {"from_address": "mail@example.com"}
    rows = rows_from_jira_brief(brief, overrides, "tata")
    batch = preview(rows)
    assert batch["ready"] == 0
    assert batch["blocked"] == 1
    issues = batch["items"][0]["issues"]
    # Must fail only on segment specification, NOT on campaign_name or schedule!
    assert "Specify exactly one of segment_id or segment_name; never all users" in issues
    assert not any("campaign_name is required" in i for i in issues)
    assert not any("scheduled_at and timezone are required" in i for i in issues)

def test_whatsapp_body_template_matching_and_placeholder_extraction(monkeypatch):
    from template_identifier import (
        extract_variables_from_template_body,
        match_whatsapp_template_and_variables,
        _CACHED_MOENGAGE_WHATSAPP_TEMPLATES,
    )
    sample_tpl = {
        "id": "6a9a6585f5b9b18f8c972b54",
        "name": "hr_btp_updated_7thsept_en",
        "body": (
            "Dear Colleague,\n \n"
            "You have been invited to attend the Behvaioural training program \n"
            "Program Name : {{1}} \n \n"
            "🗓 Session Date: {{2}}\n"
            "🕙 Time: {{3}}\n"
            "📍 Platform: MS Teams \n"
            "👉 {{4}}\n\n"
            "Attendance is mandatory. Please ensure you join on time.Non-attendance will lead to strict disciplinary action.\n\n"
            "Regards, \nL&D – Tata Capital Ltd."
        ),
        "sender_id": "685a47bfadfdfd854f960ab5",
        "provider": "Karix",
    }
    monkeypatch.setitem(_CACHED_MOENGAGE_WHATSAPP_TEMPLATES, "tata", (9999999999.0, [sample_tpl]))

    client_msg = (
        "Dear Colleague,\n \n"
        "You have been invited to attend the Behvaioural training program \n"
        "Program Name : Thrive to Amplify - Building teams that drive the Impact\n \n"
        "🗓 Session Date: 09th October 2026_TOMMORROW\n"
        "🕙 Time: 11:00 am - 01:00 pm\n"
        "📍 Platform: MS Teams \n"
        "👉 https://teams.microsoft.com/meet/43017094361739?p=test1234\n\n"
        "Attendance is mandatory. Please ensure you join on time.Non-attendance will lead to strict disciplinary action.\n\n"
        "Regards, \nL&D – Tata Capital Ltd."
    )

    match = match_whatsapp_template_and_variables(client_msg, account="tata")
    assert match is not None
    assert match["matched"] is True
    assert match["template_id"] == "6a9a6585f5b9b18f8c972b54"
    assert match["template_name"] == "hr_btp_updated_7thsept_en"
    assert "Thrive to Amplify" in match["placeholders"]["{{1}}"]
    assert "09th October 2026" in match["placeholders"]["{{2}}"]
    assert "11:00 am - 01:00 pm" in match["placeholders"]["{{3}}"]
    assert "https://teams.microsoft.com/meet" in match["placeholders"]["{{4}}"]

    # Test via rows_from_file & prepare_batch
    csv_content = (
        "channel,campaign_name,Message Body,Date of Trigger,Time of Trigger\n"
        f"WHATSAPP,Test_WA_Campaign,\"{client_msg}\",2026-10-09,11:00 AM\n"
    )
    rows = rows_from_file("wa.csv", csv_content.encode("utf-8"), "tata")
    assert len(rows) == 1
    row0 = rows[0]
    assert row0["whatsapp_template_id"] == "6a9a6585f5b9b18f8c972b54"
    assert row0["matched_template_name"] == "hr_btp_updated_7thsept_en"

    catalog = {
        "account": "tata",
        "workspace_id": "WS123",
        "segments": [{"id": "65cf4af4d4c88174e5ad186e", "name": "Test_FSTP_Pranav_1602"}],
        "whatsapp_senders": [{"sender_name": "Tata Capital Financial Services Limited", "provider": "KARIX"}],
        "whatsapp_templates": [{"id": "6a9a6585f5b9b18f8c972b54", "name": "hr_btp_updated_7thsept_en"}],
    }
    batch = prepare_batch(rows, "tata", catalog, "tester@example.com", source_type="spreadsheet")
    assert batch["ready"] == 1
    assert batch["blocked"] == 0
    cand = batch["items"][0]["candidate_v5_payload"]
    wa_conf = cand["campaign_content"]["content"]["whatsapp"]
    assert wa_conf["template_id"] == "6a9a6585f5b9b18f8c972b54"
    assert wa_conf["template_name"] == "hr_btp_updated_7thsept_en"
    assert wa_conf["body_placeholders"]["{{1}}"] == match["placeholders"]["{{1}}"]

def test_google_sheets_export_with_trailing_empty_columns_and_auxiliary_sheets():
    from openpyxl import Workbook
    import io

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Campaigns"
    # Simulate Google Sheets row 1 with 7 headers + 3 trailing empty cells (columns H, I, J)
    ws1.append(["Sr. No", "Message Body", "File Name on FileZilla", "Emp Count", "Date of Trigger", "Time of Trigger", "Teams Link", None, None, ""])
    # Simulate data row with 7 values + 3 trailing empty cells
    ws1.append(["1", "Hello message body", "FileZilla_Batch", "100", "2026-10-09", "11:00 AM", "https://teams.microsoft.com/meet/123", None, None, ""])
    # Add an empty auxiliary sheet (like Sheet2)
    ws2 = wb.create_sheet(title="EmptySheet")

    bio = io.BytesIO()
    wb.save(bio)
    data = bio.getvalue()

    # Must succeed without raising "Spreadsheet headers must be unique, non-empty"
    parsed = rows_from_file("google_sheets_export.xlsx", data, "tata")
    assert len(parsed) == 1
    row = parsed[0]
    assert row["Sr. No"] == "1"
    assert row["Message Body"] == "Hello message body"
    assert row["File Name on FileZilla"] == "FileZilla_Batch"
    assert row["Emp Count"] == "100"
    assert row["Date of Trigger"] == "2026-10-09"
    assert row["Time of Trigger"] == "11:00 AM"
    assert row["Teams Link"] == "https://teams.microsoft.com/meet/123"
    # Verify trailing empty headers were stripped and not assigned
    assert None not in row
    assert "" not in row

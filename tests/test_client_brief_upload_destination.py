"""Client briefs filter channels before enforcing the selected upload destination."""

import openpyxl
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

import api


HEADERS = {
    3: "Template Whitelisted\nY/N*",
    4: "If Template Whitelisted - Template Name/ID*",
    5: "Template Category*",
    6: "Language*",
    7: "Header Type*",
    8: "Header Text/Card Title*",
    9: 'Body/Card *\n(For Carousel, Please mention "Card 1", "Card 2", ... Clearly to avoid Confusion)',
    10: "Footer",
    11: "Button Type*",
    12: 'Button Text*\n(For Carousel, Please mention "Button 1", "Button 2", ... Clearly to avoid Confusion)',
    13: "Redirection Link *",
    14: "Bot_ID/Header Type (SMS)",
    15: "Image Link",
    17: "account",
    18: "client",
    19: "waba_id",
}


def brief_path(tmp_path, overrides=None):
    """Use the real two-level Brief layout, with an active instruction worksheet."""
    book = openpyxl.Workbook()
    instruction = book.active
    instruction.title = "Sheet1"
    instruction.append(["CHANNEL", "CAMPAIGN NAME", "CONTENT", "IMAGE"])
    instruction.append(["Whatsapp", None, "Title/Sub title:\nHeader:\nBody:\nCTA:\nurl link:"])
    instruction.append(["RCS", None, None])
    sheet = book.create_sheet("Brief")
    sheet.merge_cells("A6:A8")
    sheet["A6"] = "Content Type*"
    sheet.merge_cells("B6:B8")
    sheet["B6"] = "Channel*"
    sheet.merge_cells("C6:O7")
    sheet["C6"] = "Content"
    for column, header in HEADERS.items():
        sheet.cell(8, column, header)
    for number, channel in ((9, "WA"), (10, "Whatsapp"), (11, "RCS"), (12, "SMS"), (17, "RCS")):
        values = {
            1: "Retargeting", 2: channel, 3: "No", 5: "PROMOTIONAL", 6: "en",
            7: "Text", 9: f"Original customer copy from row {number}.",
            14: "selected-bot" if channel == "RCS" else None,
            17: "tata", 18: "tata", 19: "selected-waba" if channel in {"WA", "Whatsapp"} else None,
        }
        values.update((overrides or {}).get(number, {}))
        for column, value in values.items():
            sheet.cell(number, column, value)
    path = tmp_path / "client-brief.xlsx"
    book.save(path)
    book.close()
    return path


@pytest.fixture(autouse=True)
def selected_destination(monkeypatch):
    monkeypatch.setattr(api, "get_rcs_bot_id", lambda account: "selected-bot")
    monkeypatch.setattr(api, "get_waba_id", lambda account: "selected-waba")


@pytest.mark.parametrize("channel,selected_rows", [
    ("rcs", {11, 17}), ("whatsapp", {9, 10}), ("sms", {12}),
])
def test_mixed_brief_filters_other_channels_before_ownership(tmp_path, channel, selected_rows):
    overrides = {
        number: {14: "other-bot", 17: "other-account", 18: "other-client", 19: "other-waba"}
        for number in {9, 10, 11, 12, 17} - selected_rows
    }
    path = brief_path(tmp_path, overrides)
    api._validate_upload_destination(str(path), ".xlsx", "tata", channel)


@pytest.mark.parametrize("channel,row", [("rcs", 17), ("whatsapp", 10), ("sms", 12)])
@pytest.mark.parametrize("column", [17, 18])
def test_selected_brief_channel_cannot_override_account(tmp_path, channel, row, column):
    path = brief_path(tmp_path, {row: {column: "another-account"}})
    with pytest.raises(HTTPException) as error:
        api._validate_upload_destination(str(path), ".xlsx", "tata", channel)
    assert error.value.status_code == 400
    assert error.value.detail == f"Row {row}: account does not match the selected destination."


def test_selected_whatsapp_alias_still_checks_waba(tmp_path):
    path = brief_path(tmp_path, {10: {19: "another-waba"}})
    with pytest.raises(HTTPException) as error:
        api._validate_upload_destination(str(path), ".xlsx", "tata", "whatsapp")
    assert error.value.status_code == 400
    assert error.value.detail == "Row 10: WABA does not match the selected account."


@pytest.mark.parametrize("endpoint", ["preview", "submit"])
def test_upload_rejects_explicit_row17_bot_on_same_account(
    tmp_path, monkeypatch, provision_user, endpoint,
):
    monkeypatch.setattr(api, "get_rcs_bot_ids", lambda account: ["selected-bot", "other-bot"])
    path = brief_path(tmp_path, {17: {14: "other-bot"}})
    _, headers = provision_user("tata")
    response = TestClient(api.app).post(f"/api/{endpoint}", params={
        "account": "tata", "channel": "rcs", "expected_account": "tata", "expected_channel": "rcs",
    }, headers=headers, files={"file": (path.name, path.read_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
    assert response.status_code == 400, response.text
    assert response.json()["detail"] == "Row 17: RCS bot ID does not match the selected account."


@pytest.mark.parametrize("field,value,reason", [
    ("channel", "whatsapp", "channel does not match the selected destination"),
    ("account", "another-account", "account does not match the selected destination"),
])
def test_unrecognized_workbook_retains_flat_destination_checks(tmp_path, field, value, reason):
    book = openpyxl.Workbook()
    book.active.append(["template_name", "body_text", field])
    book.active.append(["original", "Original flat copy.", value])
    path = tmp_path / "flat.xlsx"
    book.save(path)
    book.close()
    with pytest.raises(HTTPException) as error:
        api._validate_upload_destination(str(path), ".xlsx", "tata", "rcs")
    assert error.value.status_code == 400
    assert error.value.detail == f"Row 2: {reason}."


def test_unrecognized_sms_workbook_retains_dlt_validation(tmp_path):
    book = openpyxl.Workbook()
    book.active.append(["template_name", "body_text", "account"])
    book.active.append(["original", "Original flat copy.", "another-account"])
    path = tmp_path / "flat-sms.xlsx"
    book.save(path)
    book.close()
    api._validate_upload_destination(str(path), ".xlsx", "tata", "sms")

"""
Unit tests for Excel spreadsheet parser enhancements in briefing_parser.py.
Verifies:
1. Channel sheet name matching handles variations ("WhatsApp Content", "RCS Copies", "SMS_1").
2. Grid parser handles single empty rows as paragraph separators without prematurely splitting templates.
3. Grid parser correctly splits on true section dividers (channel headers, multiple empty rows).
"""

import openpyxl

from briefing_parser import (
    _match_sheet_channel,
    _parse_excel_channel_sheets,
    _parse_excel_grid_messages,
)


def test_match_sheet_channel_variations():
    """Verify sheet channel matcher recognizes common client naming variations."""
    assert _match_sheet_channel("WhatsApp Content") == "WA"
    assert _match_sheet_channel("WhatsApp Copies") == "WA"
    assert _match_sheet_channel("WA_1") == "WA"
    assert _match_sheet_channel("WA Content") == "WA"
    assert _match_sheet_channel("WHATSAPP") == "WA"
    assert _match_sheet_channel("WA") == "WA"

    assert _match_sheet_channel("RCS Copies") == "RCS"
    assert _match_sheet_channel("RCS Content") == "RCS"
    assert _match_sheet_channel("RCS_1") == "RCS"
    assert _match_sheet_channel("RCS") == "RCS"

    assert _match_sheet_channel("SMS_1") == "SMS"
    assert _match_sheet_channel("SMS Copies") == "SMS"
    assert _match_sheet_channel("SMS Content") == "SMS"
    assert _match_sheet_channel("SMS") == "SMS"

    # Non-channel sheets should return None
    assert _match_sheet_channel("Summary") is None
    assert _match_sheet_channel("Email Mailer") is None
    assert _match_sheet_channel("Warning Notes") is None
    assert _match_sheet_channel("Sheet1") is None


def test_parse_excel_channel_sheets_with_varied_names():
    """Verify _parse_excel_channel_sheets extracts templates from varied sheet names."""
    wb = openpyxl.Workbook()
    ws_wa = wb.active
    ws_wa.title = "WhatsApp Content"

    ws_wa.append(["template_name", "body", "header_text", "button_text", "button_url"])
    ws_wa.append(
        [
            "diwali_offer_wa",
            "Dear Customer, enjoy special Diwali loan offers with low interest rates from Tata Capital.",
            "Festive Offer",
            "Apply Now",
            "https://www.tatacapital.com/diwali",
        ]
    )

    ws_rcs = wb.create_sheet(title="RCS Copies")
    ws_rcs.append(["Copy Header", "Message Body"])
    ws_rcs.append(
        [
            "RCS Body",
            "Title: Dream Car Loans\nBody: Get instant car loan approval at 8.75% p.a.\nCTA: Check Eligibility",
        ]
    )

    ws_sms = wb.create_sheet(title="SMS_1")
    ws_sms.append(["SMS Text", "Dear Customer, your EMI of Rs. 15,000 is due on 5th Oct. Pay now: https://tcl.in"])

    items = _parse_excel_channel_sheets(wb)
    channels = {item["channel"] for item in items}
    assert "WA" in channels
    assert "RCS" in channels
    assert "SMS" in channels

    wa_item = next(i for i in items if i["channel"] == "WA")
    assert "special Diwali loan offers" in wa_item["text"]
    assert wa_item.get("header") == "Festive Offer"
    assert wa_item.get("button_text") == "Apply Now"

    rcs_item = next(i for i in items if i["channel"] == "RCS")
    assert "instant car loan approval" in rcs_item["text"]
    assert rcs_item.get("title") == "Dream Car Loans"


def test_parse_excel_grid_messages_does_not_split_on_blank_row():
    """Verify _parse_excel_grid_messages preserves multi-paragraph copy with blank rows."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "LAP Content"

    # Multi-row grid layout where Row 0 has channel, Row 1 has paragraph 1,
    # Row 2 is an empty spacing row, Row 3 has paragraph 2, Row 4 has CTA
    ws.append(["Channel", "Variant 1"])
    ws.append(["SMS", "Dear Customer, get pre-approved Home Loan up to Rs. 75 Lakhs."])
    ws.append(["", ""])  # Single empty row inside body for paragraph spacing
    ws.append(["", "Special festival interest rate of 8.5% p.a. valid till 31st Oct."])
    ws.append(["", "Apply online: https://www.tatacapital.com/homeloan"])
    ws.append(["", ""])
    ws.append(["", ""])  # Two consecutive empty rows -> signals section end

    # Second channel section
    ws.append(["RCS", "Upgrade to a luxurious apartment with Tata Capital Home Loans."])
    ws.append(["", "Fast approvals with minimal documentation."])

    items = _parse_excel_grid_messages(wb)
    assert len(items) == 2, f"Expected exactly 2 templates (1 SMS, 1 RCS), got {len(items)}"

    sms_item = next(i for i in items if i["channel"] == "SMS")
    assert "Dear Customer, get pre-approved" in sms_item["text"]
    assert "Special festival interest rate" in sms_item["text"]
    assert "Apply online:" in sms_item["text"]

    rcs_item = next(i for i in items if i["channel"] == "RCS")
    assert "Upgrade to a luxurious apartment" in rcs_item["text"]
    assert "Fast approvals" in rcs_item["text"]

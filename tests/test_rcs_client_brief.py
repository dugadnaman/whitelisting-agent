"""Structured intake regressions; no external provider calls are permitted."""
import io

import openpyxl
import pytest
from openpyxl.drawing.image import Image as ExcelImage
from PIL import Image

import rcs_config
import rcs_loader
from client_brief import read_client_brief_rows


HEADERS = {
    3: 'Template Whitelisted\nY/N*',
    4: 'If Template Whitelisted - Template Name/ID*',
    5: 'Template Category*', 6: 'Language*', 7: 'Header Type*',
    8: 'Header Text/Card Title*',
    9: 'Body/Card *\n(For Carousel, Please mention "Card 1", "Card 2", ... Clearly to avoid Confusion)',
    10: 'Footer', 11: 'Button Type*',
    12: 'Button Text*\n(For Carousel, Please mention "Button 1", "Button 2", ... Clearly to avoid Confusion)',
    13: 'Redirection Link *', 14: 'Bot_ID/Header Type (SMS)', 15: 'Image Link',
}


@pytest.fixture(autouse=True)
def isolated(monkeypatch, tmp_path):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(rcs_config, '_load_env_file', lambda: None)
    monkeypatch.setenv('TCL_PROMO_RCS_BOT_ID', 'explicit-bot')
    monkeypatch.setattr(rcs_loader, 'get_rcs_bot_id', lambda client: 'explicit-bot')
    monkeypatch.setattr(rcs_loader, 'get_rcs_entity_id', lambda client: 'entity')
    import rcs_client
    monkeypatch.setattr(rcs_client, 'upload_rcs_media', lambda *args, **kwargs: pytest.fail('Provider upload forbidden'))


def workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.title = 'Sheet1'
    wb.active.append(['CHANNEL', 'CAMPAIGN NAME', 'CONTENT', 'IMAGE'])
    wb.active.append(['RCS', 'Instruction, not copy', 'Title/Sub title:', 'Image:'])
    ws = wb.create_sheet('Brief')
    ws.merge_cells('A6:A8')
    ws.merge_cells('B6:B8')
    ws.merge_cells('C6:O7')
    ws['A6'] = 'Content Type*'
    ws['B6'] = 'Channel*'
    ws['C6'] = 'Content'
    for col, title in HEADERS.items():
        ws.cell(8, col, title)
    return wb, ws, tmp_path / 'BRIEF.xlsx'


def row(ws, number, kind='Richcard', bot='explicit-bot'):
    values = {1: 'Retargeting', 2: 'RCS', 3: 'No', 5: 'PROMOTIONAL', 6: 'en',
              7: kind, 8: 'Actual title', 9: 'Actual body\n\nSecond paragraph',
              11: 'Website (RCS)', 12: 'Know More', 13: 'https://example.test/offer', 14: bot}
    if kind == 'Carousel':
        values.update({8: 'Card 1 :- First\nCard 2 :- Second',
                       9: 'Card 1 :- First body\ninner line\n\nparagraph\nCard 2 :- Second body',
                       12: 'Button 1 :- First action\nButton 2 :- Second action'})
    for col, value in values.items():
        ws.cell(number, col, value)


def add_image(ws, cell, color='red'):
    stream = io.BytesIO()
    Image.new('RGB', (160, 90), color).save(stream, format='PNG')
    stream.seek(0)
    ws.add_image(ExcelImage(stream), cell)


def load(path):
    return rcs_loader.load_rcs_from_excel(str(path), client='tcl_promo', upload_media=False)


def test_merged_headers_non_active_sheet_richcard_and_provenance(tmp_path):
    wb, ws, path = workbook(tmp_path)
    row(ws, 9)
    ws['B9'] = 'WA'
    row(ws, 16)
    ws['Q8'], ws['R8'], ws['S8'] = 'account', 'client', 'waba_id'
    ws['Q16'], ws['R16'], ws['S16'] = 'tcl_promo', 'tcl_promo', 'explicit-waba'
    add_image(ws, 'O16')
    wb.save(path)
    records = read_client_brief_rows(str(path))
    assert [record['source_ref'] for record in records] == ['Brief!9', 'Brief!16']
    assert records[1]['waba_id'] == 'explicit-waba'
    assert records[1]['body'] == 'Actual body\n\nSecond paragraph'
    first = load(path)
    second = load(path)
    assert len(first) == 1
    sub = first[0]
    assert sub.template_type == 'richcard'
    assert sub.card_title == 'Actual title'
    assert sub.card_description == records[1]['body']
    assert sub.source_ref == 'Brief!16'
    assert sub._excel_sheet == 'Brief' and sub._excel_row == 16
    assert sub.image_bytes
    assert sub.bot_id == 'explicit-bot'
    assert sub.template_name == second[0].template_name and len(sub.template_name) <= 25


def test_carousel_copy_shared_url_and_no_cross_row_or_sheet_media(tmp_path):
    wb, ws, path = workbook(tmp_path)
    row(ws, 11, 'Carousel')
    row(ws, 13, 'Carousel')
    row(ws, 16)
    add_image(ws, 'O13')  # Incomplete carousel set cannot identify a particular card.
    add_image(ws, 'O16', 'blue')
    add_image(wb['Sheet1'], 'D11', 'green')
    wb.save(path)
    subs = load(path)
    assert len({sub.template_name for sub in subs}) == 3
    for sub in subs[:2]:
        assert len(sub.carousel_cards) == 2
        assert sub.carousel_cards[0]['cardDescription'] == 'First body\ninner line\n\nparagraph'
        assert [card['suggestions'][0]['text'] for card in sub.carousel_cards] == ['First action', 'Second action']
        assert all(card['suggestions'][0]['url'].startswith('https://example.test/offer') for card in sub.carousel_cards)
        assert all(not card.get('mediaUrl') and not card.get('image_bytes') for card in sub.carousel_cards)
    assert subs[2].image_bytes


@pytest.mark.parametrize('layout', ['nested', 'slash'])
def test_multiple_buttons_per_card_and_parallel_mappings(tmp_path, layout):
    wb, ws, path = workbook(tmp_path)
    row(ws, 11, 'Carousel')
    if layout == 'nested':
        ws['L11'] = 'Card 1 :-\nButton 1 :- One\nButton 2 :- Two\nCard 2 :-\nButton 1 :- Three\nButton 2 :- Four'
        ws['M11'] = 'Card 1 :-\nButton 1 :- https://one.test/\nButton 2 :- https://two.test/\nCard 2 :-\nButton 1 :- https://three.test/\nButton 2 :- https://four.test/'
    else:
        ws['L11'] = 'Card 1 / Button 1 :- One\nCard 1 / Button 2 :- Two\nCard 2 / Button 1 :- Three\nCard 2 / Button 2 :- Four'
        ws['M11'] = 'https://shared.test/'
    wb.save(path)
    cards = load(path)[0].carousel_cards
    assert [[button['text'] for button in card['suggestions']] for card in cards] == [['One', 'Two'], ['Three', 'Four']]
    if layout == 'nested':
        assert cards[1]['suggestions'][1]['url'].startswith('https://four.test/')
    else:
        assert all(button['url'].startswith('https://shared.test/') for card in cards for button in card['suggestions'])


@pytest.mark.parametrize('column,value', [
    ('I11', 'Card 1 :- Only one'),
    ('L11', 'Button 1 :- Only one'),
    ('L11', 'Unassigned single button'),
    ('L11', 'Card 1 / Button 1 :- One\nCard 1 / Button 1 :- Duplicate\nCard 2 / Button 1 :- Two'),
    ('M11', 'Card 1 :-\nButton 1 :- https://one.test/\nButton 2 :- https://two.test/\nCard 2 :- https://three.test/'),
    ('O11', 'https://single-ambiguous-image.test/image.png'),
    ('L11', ''),
])
def test_contradictory_counts_and_ambiguous_mapping_are_rejected(tmp_path, column, value):
    wb, ws, path = workbook(tmp_path)
    row(ws, 11, 'Carousel')
    ws[column] = value
    wb.save(path)
    with pytest.raises(ValueError, match='Brief!11:'):
        load(path)


def test_explicit_other_bot_is_never_overridden_or_uploaded(tmp_path):
    wb, ws, path = workbook(tmp_path)
    row(ws, 16)
    row(ws, 17, bot='other-explicit-bot')
    ws['F17'] = 'tel'
    add_image(ws, 'O16')
    wb.save(path)
    records = read_client_brief_rows(str(path))
    assert records[-1]['bot_id'] == 'other-explicit-bot'
    with pytest.raises(ValueError, match='Brief!17: RCS row bot ID'):
        rcs_loader.load_rcs_from_excel(str(path), client='tcl_promo', upload_media=True)


def test_flat_workbook_is_not_recognized_as_structured_brief(tmp_path):
    path = tmp_path / 'flat.xlsx'
    wb = openpyxl.Workbook()
    wb.active.append(['template_name', 'body', 'bot_id'])
    wb.active.append(['existing_flat', 'Existing flat copy', 'explicit-bot'])
    wb.save(path)
    assert read_client_brief_rows(str(path)) is None
    assert load(path)[0].template_name == 'existing_flat'


def test_actual_style_four_card_english_and_hindi_rows(tmp_path, caplog):
    wb, ws, path = workbook(tmp_path)
    for number, language in ((11, ''), (13, 'Hin')):
        row(ws, number, 'Carousel')
        ws.cell(number, 6, language)
        ws.cell(number, 8, '\n'.join(f'Card {card} :- Title {card}' for card in range(1, 5)))
        ws.cell(number, 9, '\n'.join(f'Card {card} :- Body {card}\n\nInner paragraph' for card in range(1, 5)))
        ws.cell(number, 12, '\n'.join(f'Button {card} :- Action {card}' for card in range(1, 5)))
    add_image(ws, 'O13')
    wb.save(path)
    records = read_client_brief_rows(str(path))
    assert [record['language'] for record in records] == ['', 'Hin']
    subs = load(path)
    assert [len(sub.carousel_cards) for sub in subs] == [4, 4]
    assert [card['suggestions'][0]['text'] for card in subs[1].carousel_cards] == ['Action 1', 'Action 2', 'Action 3', 'Action 4']
    assert all(not card.get('mediaUrl') for sub in subs for card in sub.carousel_cards)
    assert 'Brief!13: 1 embedded images cannot map unambiguously to 4 media targets' in caplog.text

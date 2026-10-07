"""Deterministic worker regressions. No real accounts, Google RPCs or campaigns.

Run from Backend with python -m unittest discover -s tests -p 'test_*.py'.
The private key below is generated afresh for validation tests, never a deployed key.
"""
import asyncio
import json
import os
import subprocess
import tempfile
import threading
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.utils.cell import range_boundaries
from pydantic import ValidationError

# Never read operator/deployment credentials or persistent state while importing tests.
with tempfile.TemporaryDirectory() as import_directory, patch.dict(os.environ, {
    "STORAGE_DIR": import_directory, "GOOGLE_SERVICE_ACCOUNT_JSON": "",
    "GOOGLE_SERVICE_ACCOUNT_FILE": "", "GOOGLE_SPREADSHEET_URL": "",
    "GOOGLE_WORKSHEET_NAME": "Mastersheet", "MOENGAGE_UI_CONFIG_JSON": "{}",
    "MOENGAGE_REMOTE_CDP_URL": "", "MOENGAGE_BROWSER_LOGIN_URL": "",
    "MOENGAGE_DASHBOARD_URL": "https://dashboard.moengage.com/", "MOENGAGE_MODE": "browser",
    "ALLOW_MOCK_WRITES": "false", "APPAREL_ATTRIBUTION_TOKEN": "",
}):
    from app.api import routes
    from app.config.settings import Settings, spreadsheet_identity, strong_token, validate_private_cdp, validate_ui_config
    from app.core import dependencies
    from app.main import app
    from app.models.report import CampaignMetrics, JobState, ReportJob, RowResult, SheetConnection
    from app.schemas.report_schema import MoEngageSessionRequest, SetupRequest, StartJobRequest
    from app.services.google_sheet_service import GoogleSheetService, finish_thread_call
    from app.services.moengage_browser_service import BrowserAutomationError, MoEngageBrowserService
    from app.services.moengage_service import MoEngageService
    from app.services.report_service import ReportService
    from app.utils.excel_utils import parse_tracking_range
    from app.services.excel_service import ExcelService


TOKEN = "7aP9!xD3#uF6@qR1$sT8%vW2&yZ5*bC0"
SHEET_ID = "approved_apparel_sheet_identity_1234567890"
SHEET_URL = f"https://docs.google.com/spreadsheets/d/{SHEET_ID}/edit"
UI = {"workflow": "recorded_behavior", "query_url_map": {"Aldo": "https://dashboard.moengage.com/behavior?did=report&chartId=chart"}}


class WorksheetFixture:
    def __init__(self):
        self.sheet = Workbook().active
        self.sheet.title = "Mastersheet"
        self.headers = list(ExcelService.INPUT_HEADERS) + [
            "Campaign Purchased Customers", "Campaign Purchased Customers - Online",
            "Campaign Purchased Customers - Offline", "Influenced Revenue",
            "Influenced Revenue - Online", "Influenced Revenue - Offline",
        ]
        self.sheet.append(self.headers)
        self.sheet.append(["01/08/2026", "WhatsApp", "Aldo", "Completed", "Online", "1 Aug", "complete", 7, 7, "", 450, 450, ""])
        self.sheet.append(["01/08/2026", "SMS", "Aldo", "Incomplete", "Overall", "1 Aug", "incomplete", "", "", "", "", "", ""])
        self.sheet["AA2"] = "unrelated customer data"
        self.writes = []

    def get(self, cell_range, **_kwargs):
        if cell_range == "A:AF":
            return list(self.sheet.iter_rows(max_col=32, values_only=True))
        left, top, right, bottom = range_boundaries(cell_range)
        return list(self.sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right, values_only=True))

    def batch_update(self, updates, **_kwargs):
        self.writes.append(updates)
        for update in updates:
            self.sheet[update["range"]] = update["values"][0][0]


class WorkerRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Fresh synthetic RSA key validates the real google-auth parser without a network call.
        result = subprocess.run([
            "openssl", "genpkey", "-algorithm", "RSA", "-pkeyopt", "rsa_keygen_bits:2048",
        ], check=True, capture_output=True)
        cls.key = json.dumps({
            "type": "service_account", "project_id": "synthetic-test-project",
            "client_email": "test@synthetic-test-project.iam.gserviceaccount.com",
            "private_key": result.stdout.decode(), "token_uri": "https://oauth2.googleapis.com/token",
        }).encode()

    def setUp(self):
        self.folder = tempfile.TemporaryDirectory()
        self.directory = Path(self.folder.name)
        self.settings = Settings(
            storage_dir=self.directory, machine_token=TOKEN,
            google_service_account_file=self.directory / "google-service-account.json",
            google_spreadsheet_url=SHEET_URL, moengage_ui_config=json.loads(json.dumps(UI)),
            moengage_remote_cdp_url="http://browser.internal:9222",
            moengage_browser_login_url="https://login.corporate.test/",
        )
        self.service = ReportService(self.settings)
        self.service.google.install_credentials(self.key)
        self.worksheet = WorksheetFixture()
        self.service.google._worksheet_cache[(SHEET_ID, "Mastersheet")] = self.worksheet
        rows, _ = self.service.google._read_worksheet(self.worksheet)
        self.connection = SheetConnection(
            "connection", SHEET_ID, SHEET_URL, "Synthetic sheet", "Mastersheet", len(rows),
            ["Aldo"], ["WhatsApp", "SMS"], ["Online", "Overall"], [date(2026, 8, 1)], [], campaigns=rows,
        )
        self.service.sheet_connections[self.connection.id] = self.connection
        app.dependency_overrides[dependencies.get_report_service] = lambda: self.service
        self.token_patch = patch.object(dependencies.settings, "machine_token", TOKEN)
        self.token_patch.start()
        self.client = TestClient(app)
        self.headers = {"X-Apparel-Attribution-Token": TOKEN, "X-Apparel-Actor": "verified-id|operator@corporate.test"}

    def tearDown(self):
        self.client.close()
        app.dependency_overrides.clear()
        self.token_patch.stop()
        self.folder.cleanup()

    def test_private_routes_require_exact_machine_token_and_public_health_is_minimal(self):
        self.assertEqual(self.client.get("/healthz").json(), {"status": "ok"})
        for path in ["/api/health", "/api/google/config", "/api/setup", "/api/jobs", "/api/moengage/session"]:
            self.assertEqual(self.client.get(path).status_code, 401)
            self.assertEqual(self.client.get(path, headers={"X-Apparel-Attribution-Token": TOKEN + "wrong"}).status_code, 401)
        self.assertEqual(self.client.get("/api/google/config", headers=self.headers).status_code, 200)
        for path in ["/", "/docs", "/redoc", "/openapi.json", "/api/upload", "/assets/anything"]:
            self.assertEqual(self.client.get(path, headers=self.headers).status_code, 404)
        with patch.object(dependencies.settings, "machine_token", ""):
            self.assertEqual(self.client.get("/api/google/config", headers=self.headers).status_code, 503)
        self.assertFalse(strong_token("a" * 40))
        self.assertFalse(strong_token("0123456789abcdef" * 3))
        self.assertTrue(strong_token(TOKEN))

    def test_verified_actor_and_safe_config_no_password_or_upload_routes(self):
        unaudited = self.client.put("/api/setup", json={"spreadsheet_url": SHEET_URL, "worksheet_name": "Mastersheet", "ui_config": UI}, headers={"X-Apparel-Attribution-Token": TOKEN})
        self.assertEqual(unaudited.status_code, 400)
        for payload in [{"profile_id": "default", "password": "never-forward"}, {"profile_id": "invalid-profile"}]:
            self.assertEqual(self.client.post("/api/moengage/session/start", headers=self.headers, json=payload).status_code, 422)
        with self.assertRaises(ValidationError):
            StartJobRequest(upload_id="workbook", brands=["Aldo"], channels=["SMS"], sent_date="2026-08-01")
        with self.assertRaises(ValidationError):
            StartJobRequest(sheet_connection_id="connection", brands=[" "], channels=["SMS"], sent_date="2026-08-01")
        with self.assertRaises(ValidationError):
            StartJobRequest(sheet_connection_id="connection", brands=["Aldo"], channels=["SMS"], sent_date_from="2025-01-01", sent_date_to="2026-08-01")

    def test_setup_persists_restart_and_upload_does_not_override_sealed_key(self):
        changed_url = SHEET_URL.replace(SHEET_ID, "updated_approved_identity_123456789012345")
        response = self.client.put("/api/setup", headers=self.headers, json={
            "spreadsheet_url": changed_url, "worksheet_name": "Approved Apparel", "ui_config": UI,
        })
        self.assertEqual(response.status_code, 200)
        self.assertNotIn("private_key", response.text)
        self.assertEqual(self.settings.setup_path.stat().st_mode & 0o777, 0o600)
        with patch.dict(os.environ, {
            "STORAGE_DIR": str(self.directory), "GOOGLE_SERVICE_ACCOUNT_JSON": "",
            "GOOGLE_SERVICE_ACCOUNT_FILE": "", "GOOGLE_SPREADSHEET_URL": SHEET_URL,
            "GOOGLE_WORKSHEET_NAME": "Old value", "MOENGAGE_UI_CONFIG_JSON": "{}",
            "MOENGAGE_REMOTE_CDP_URL": "", "MOENGAGE_BROWSER_LOGIN_URL": "",
            "MOENGAGE_DASHBOARD_URL": "", "MOENGAGE_MODE": "browser", "ALLOW_MOCK_WRITES": "false",
        }):
            restored = Settings.from_env()
        self.assertEqual(restored.google_spreadsheet_url, changed_url)
        self.assertEqual(restored.google_worksheet_name, "Approved Apparel")
        self.assertEqual(restored.moengage_ui_config, UI)
        restored_service = ReportService(restored)
        self.assertTrue(restored_service.google.configured)
        self.assertEqual(restored_service.google.service_account_email(), "test@synthetic-test-project.iam.gserviceaccount.com")
        self.assertEqual(restored.google_service_account_file.stat().st_mode & 0o777, 0o600)
        self.settings.sealed_google_credentials = True
        before = self.settings.google_service_account_file.read_bytes()
        rejected = self.client.post("/api/google/credentials", headers=self.headers, files={"credential": ("credential.json", self.key, "application/json")})
        self.assertEqual(rejected.status_code, 409)
        self.assertEqual(self.settings.google_service_account_file.read_bytes(), before)

    def test_invalid_config_and_key_fail_closed_and_approved_sheet_cannot_switch(self):
        for config in [{"mode": "mock"}, {"workflow": "recorded_behavior", "query_url_map": {"Aldo": "http://localhost/private"}}, {"workflow": "api"}]:
            response = self.client.put("/api/setup", headers=self.headers, json={"spreadsheet_url": SHEET_URL, "worksheet_name": "Mastersheet", "ui_config": config})
            self.assertEqual(response.status_code, 422)
        self.assertEqual(self.client.post("/api/google/connect", headers=self.headers, json={"spreadsheet_url": "https://evil.test/spreadsheets/d/" + SHEET_ID, "worksheet_name": "Mastersheet"}).status_code, 422)
        self.assertEqual(self.client.post("/api/google/connect", headers=self.headers, json={"spreadsheet_url": SHEET_URL.replace(SHEET_ID, "another_sheet_identity_1234567890"), "worksheet_name": "Mastersheet"}).status_code, 403)
        self.assertEqual(self.client.post("/api/google/connect", headers=self.headers, json={"spreadsheet_url": SHEET_URL, "worksheet_name": "Other"}).status_code, 403)
        previous = self.settings.google_service_account_file.read_bytes()
        key = json.loads(self.key)
        key["token_uri"] = "http://169.254.169.254/secret"
        response = self.client.post("/api/google/credentials", headers=self.headers, files={"credential": ("credential.json", json.dumps(key), "application/json")})
        self.assertEqual(response.status_code, 422)
        self.assertEqual(previous, self.settings.google_service_account_file.read_bytes())
        invalid = GoogleSheetService(self.directory / "missing.json", '{"private_key":"invalid"}')
        self.assertFalse(invalid.configured)
        for endpoint in ["http://8.8.8.8:9222", "http://169.254.169.254:9222", "https://browser.internal:9222", "http://browser.internal:9222/anything"]:
            with self.assertRaises(ValueError):
                validate_private_cdp(endpoint)

    def test_preview_parsing_warnings_follow_requested_date_range(self):
        self.worksheet.sheet.append(["01/08/2026", "SMS", "Aldo", "", "Online", "1 Aug", "bad-august"])
        self.worksheet.sheet.append(["01/07/2026", "SMS", "Aldo", "", "Online", "1 Jul", "bad-july"])
        response = self.client.get("/api/google/connections/connection/campaigns", headers=self.headers, params={"sent_date_from": "2026-07-01", "sent_date_to": "2026-07-31"})
        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["warning_sent_date_from"], "2026-07-01")
        self.assertEqual(body["warning_sent_date_to"], "2026-07-31")
        self.assertEqual(body["row_count"], 0)
        self.assertEqual(body["warnings"], ["Row 5: blank required cell(s): Campaign Name"])
        august = self.client.get("/api/google/connections/connection/campaigns", headers=self.headers, params={"sent_date_from": "2026-08-01", "sent_date_to": "2026-08-01"}).json()
        self.assertEqual(august["warnings"], ["Row 4: blank required cell(s): Campaign Name"])

    def test_job_admission_is_atomic_across_browser_authentication(self):
        async def scenario():
            authentication_started = asyncio.Event()
            release_authentication = asyncio.Event()
            hold_processing = asyncio.Event()

            async def authenticate():
                authentication_started.set()
                await release_authentication.wait()

            async def read_campaigns(*_arguments):
                await hold_processing.wait()
                return self.connection.campaigns, []

            self.service.moengage.ensure_authenticated = authenticate
            self.service.google.read_campaigns = read_campaigns
            request = StartJobRequest(sheet_connection_id="connection", brands=["Aldo"], channels=["WhatsApp"], sent_date="2026-08-01")
            first = asyncio.create_task(routes.start_job(request, self.service))
            await authentication_started.wait()
            second = asyncio.create_task(routes.start_job(request, self.service))
            await asyncio.sleep(0)
            self.assertEqual(len(self.service.jobs), 0)
            release_authentication.set()
            admitted = await first
            with self.assertRaises(HTTPException) as rejected:
                await second
            self.assertEqual(rejected.exception.status_code, 409)
            self.assertEqual(len(self.service.jobs), 1)
            self.assertTrue(self.service.has_active_work())
            await self.service.cancel_job(admitted.job_id)
            self.assertFalse(self.service.has_active_work())
        asyncio.run(scenario())

    def test_cancel_finishes_google_rpc_before_new_admission(self):
        async def scenario():
            started = threading.Event()
            release = threading.Event()
            finished = threading.Event()

            def actual_rpc():
                started.set()
                release.wait(timeout=5)
                finished.set()

            task = asyncio.create_task(finish_thread_call(actual_rpc))
            self.service.jobs["running"] = ReportJob("running", "connection", "Synthetic", state=JobState.PROCESSING)
            self.service.tasks["running"] = task
            while not started.is_set():
                await asyncio.sleep(0)
            cancellation = asyncio.create_task(routes.cancel_job("running", self.service))
            await asyncio.sleep(0)
            await asyncio.sleep(0)
            self.assertFalse(cancellation.done())
            self.assertTrue(self.service.has_active_work())
            self.assertTrue(self.service.admission_lock.locked())
            release.set()
            response = await cancellation
            self.assertTrue(finished.is_set())
            self.assertEqual(response.status, "cancelled")
            self.assertFalse(self.service.has_active_work())
        asyncio.run(scenario())

    def test_completed_rows_and_unrelated_cells_preserved_and_retry_only_failed_rows(self):
        async def scenario():
            async def authenticated():
                return None

            async def fixture_metrics(row):
                self.assertEqual(row.excel_row, 3)
                return CampaignMetrics(7, 450, 5, 2, 300, 150)

            self.service.moengage.ensure_authenticated = authenticated
            self.service.moengage.fetch_metrics = fixture_metrics
            job = self.service.create_sheet_job("connection", False, None, ["Aldo"], ["WhatsApp"], date(2026, 8, 1))
            await self.service.tasks[job.id]
            self.assertEqual(job.results[0].status, "skipped")
            self.assertEqual(self.worksheet.writes, [])
            self.assertEqual(self.worksheet.sheet["H2"].value, 7)
            self.assertEqual(self.worksheet.sheet["AA2"].value, "unrelated customer data")
            original = ReportJob("original", "connection", "Synthetic", state=JobState.FAILED)
            original.results = [
                RowResult(2, "Aldo", "WhatsApp", "Online", "complete", "Completed", "2026-08-01 → 2026-08-01", status="success"),
                RowResult(3, "Aldo", "SMS", "Overall", "incomplete", "Incomplete", "2026-08-01 → 2026-08-01", status="failed"),
            ]
            self.service.jobs[original.id] = original
            retry = self.service.retry_failed_sheet_job(original.id)
            await self.service.tasks[retry.id]
            self.assertEqual([result.excel_row for result in retry.results], [3])
            self.assertEqual(retry.results[0].status, "success")
            self.assertEqual(self.worksheet.sheet["H2"].value, 7)
            self.assertEqual(self.worksheet.sheet["AA2"].value, "unrelated customer data")
            self.assertEqual([self.worksheet.sheet.cell(3, column).value for column in range(8, 14)], [7, 5, 2, 450, 300, 150])
            self.assertTrue(all(update["range"].endswith("3") for write in self.worksheet.writes for update in write))
            response = self.client.get(f"/api/jobs/{retry.id}/results.csv", headers=self.headers)
            self.assertEqual(response.status_code, 200)
            self.assertIn("attachment", response.headers["content-disposition"])
            self.assertIn("Incomplete,incomplete,Aldo,SMS,Overall", response.text)
            self.assertIn(",7,5,2,450,300,150,", response.text)
        asyncio.run(scenario())

    def test_month_boundaries_currency_and_finance_policy_remain_repaired(self):
        self.assertEqual(parse_tracking_range("31 - 1 Feb", date(2026, 1, 31)), (date(2026, 1, 31), date(2026, 2, 1)))
        self.assertEqual(parse_tracking_range("31 - 1 Jan", date(2025, 12, 31)), (date(2025, 12, 31), date(2026, 1, 1)))
        for text, expected in {"1.2K": 1200, "2M": 2000000, "0.5b": 500000000, "₹4,99,279.00": 499279, "(₹2,500.50)": -2500.5, "0": 0}.items():
            self.assertEqual(MoEngageBrowserService._parse_number(text), expected)
        for text in ["12%", "not available", "2 orders", "1.2.3"]:
            with self.assertRaises(BrowserAutomationError):
                MoEngageBrowserService._parse_number(text)
        combined = MoEngageService._combine_metrics(
            MoEngageService._typed_metrics("Online", 5, 300), MoEngageService._typed_metrics("Offline", 2, 150),
        )
        self.assertEqual(combined, CampaignMetrics(7, 450, 5, 2, 300, 150))
        with self.assertRaises(Exception):
            MoEngageService._typed_metrics("Online", 0, 450)


if __name__ == "__main__":
    unittest.main()

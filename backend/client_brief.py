"""Read the client's structured, possibly multi-level Excel brief without guessing copy."""
import re


def _header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def _field(value) -> str | None:
    text = _header(value)
    prefixes = (
        ("channel", "channel"), ("templatecategory", "category"),
        ("language", "language"), ("headertype", "header_type"),
        ("headertextcardtitle", "card_title"), ("bodycard", "body"),
        ("footer", "footer"), ("buttontype", "button_type"),
        ("buttontext", "button_text"), ("redirectionlink", "button_url"),
        ("botid", "bot_id"), ("imagelink", "image_url"),
        ("iftemplatewhitelistedtemplatename", "template_name"),
        ("campaignname", "campaign_name"),
    )
    for prefix, field in prefixes:
        if text.startswith(prefix):
            return field
    return {
        "body": "body", "cardtitle": "card_title", "imageurl": "image_url",
        "client": "client", "account": "account", "wabaid": "waba_id",
        "whatsappbusinessaccountid": "waba_id", "templatename": "template_name",
    }.get(text)


def read_client_brief_rows(path: str) -> list[dict] | None:
    """Return structured worksheet records, or None for an ordinary flat workbook.

    Includes every explicit channel for callers to filter before destination checks.
    Never reads images, changes worksheets, fills down routing, or invents values.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True)
    records = []
    recognized = False
    try:
        for sheet in workbook:
            columns = None
            for cells in sheet.iter_rows():
                fields = {cell.column: _field(cell.value) for cell in cells if cell.value is not None}
                if {"card_title", "body", "button_text"}.issubset(set(fields.values())):
                    # Channel is commonly vertically merged over the two header levels.
                    for previous in sheet.iter_rows(min_row=max(1, cells[0].row - 2), max_row=cells[0].row):
                        for cell in previous:
                            if _field(cell.value) == "channel":
                                fields[cell.column] = "channel"
                    if "channel" in fields.values():
                        columns = {col: field for col, field in fields.items() if field}
                        recognized = True
                        continue
                if columns is None:
                    continue
                values = {field: sheet.cell(cells[0].row, col).value for col, field in columns.items()}
                if not str(values.get("channel") or "").strip():
                    continue
                record = {key: "" if value is None else str(value) for key, value in values.items()}
                for key in ("channel", "bot_id", "category", "language", "header_type", "client", "account", "waba_id"):
                    if key in record:
                        record[key] = record[key].strip()
                for key in ("bot_id", "category", "language", "header_type", "card_title", "body", "footer", "button_type", "button_text", "button_url", "image_url"):
                    record.setdefault(key, "")
                record.update(sheet=sheet.title, excel_row=cells[0].row, source_ref=f"{sheet.title}!{cells[0].row}")
                records.append(record)
    finally:
        workbook.close()
    return records if recognized else None

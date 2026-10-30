"""Pure, source-preserving SMS brief extraction and untrusted review validation."""

import hashlib
import re
import unicodedata
import zipfile
from collections import Counter
from pathlib import Path

MAX_EXPANDED_BYTES = 80 * 1024 * 1024
MAX_CELLS = 250_000
MAX_CANDIDATES = 5_000
MAX_TEXT = 32_767
_TOKEN = re.compile(r"\{\{[^{}\r\n]+\}\}|\{#[^{}\r\n]+#\}|<[^<>\r\n]+>|\[[^\[\]\r\n]+\]|\(\s*[xX]{2,}\s*\)")
_TYPED = re.compile(r"\{#(?:num|alp|urg|uro|cbn|eml)#\}\Z")
_KINDS = {"literal", "num", "alp", "urg", "uro", "cbn", "eml"}
_LANGUAGES = {
    "english": "English", "hindi": "Hindi", "marathi": "Marathi",
    "kannada": "Kannada", "gujarati": "Gujarati", "gujrati": "Gujarati",
    "tamil": "Tamil", "telugu": "Telugu", "bengali": "Bengali",
    "punjabi": "Punjabi", "malayalam": "Malayalam",
}
_SCRIPTS = {
    "Hindi": (0x0900, 0x097F), "Marathi": (0x0900, 0x097F),
    "Bengali": (0x0980, 0x09FF), "Punjabi": (0x0A00, 0x0A7F),
    "Gujarati": (0x0A80, 0x0AFF), "Tamil": (0x0B80, 0x0BFF),
    "Telugu": (0x0C00, 0x0C7F), "Kannada": (0x0C80, 0x0CFF),
    "Malayalam": (0x0D00, 0x0D7F),
}


def _norm(value):
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _channel(value):
    label = _norm(value)
    if len(label) > 85 or "\n" in str(value or ""):
        return ""
    if re.match(r"sms(?:\b|_)", label):
        return "sms"
    if re.match(r"(?:wa|whatsapp|rcs|pn|push notification|email)(?:\b|_)", label):
        return "other"
    return ""


def _metadata(value):
    label = _norm(value)
    return bool(re.match(
        r"^(?:cta\b|link\b|date\b|time of\b|testing\b|test mobile\b|database\b|"
        r"product\b|sub product\b|campaign tag\b|notes?\b|instructions?\b|"
        r"character(?:s| count)?\b|char count\b|use\s*case\b|remark\b|base\b)", label))


def _copy(value):
    if not isinstance(value, str):
        return False
    stripped = value.strip()
    if not 20 <= len(stripped) <= MAX_TEXT or _metadata(stripped):
        return False
    if stripped.startswith("="):
        return False
    if re.fullmatch(r"[\d\s+(),.;:/_-]+", stripped):
        return False
    if re.fullmatch(r"https?://\S+", stripped, re.I):
        return False
    if _norm(stripped) in {"mentioned in base", "mentioned in database", "as per planner"}:
        return False
    return bool(re.search(r"[^\W\d_]", stripped, re.UNICODE))


def _language(value):
    label = _norm(value)
    if label in _LANGUAGES:
        return _LANGUAGES[label]
    if "hinglish" in label or "highlish" in label:
        return "Hinglish"
    if "english" in label and len(label) < 50:
        return "English"
    return ""


def _language_warnings(text, language):
    if language not in _SCRIPTS:
        return []
    expected = _SCRIPTS[language]
    counts = Counter()
    for char in text:
        if not char.isalpha():
            continue
        for _name, (start, end) in _SCRIPTS.items():
            if start <= ord(char) <= end:
                counts[(start, end)] += 1
                break
    if counts and counts.most_common(1)[0][0] != expected:
        return [f"Declared {language} language does not match the dominant source script; verify the original cell. Copy has not been changed."]
    return []


def _column(index):
    result = ""
    while index:
        index, digit = divmod(index - 1, 26)
        result = chr(65 + digit) + result
    return result


def _read_sheets(path, filename):
    suffix = Path(filename).suffix.lower()
    if suffix not in {".xlsx", ".xls"}:
        raise ValueError("Upload an .xlsx or .xls client brief workbook")
    sheets = []
    total = 0
    try:
        if Path(path).stat().st_size > 20 * 1024 * 1024:
            raise ValueError("Client brief exceeds the 20 MiB upload limit")
        if suffix == ".xlsx":
            import openpyxl

            with zipfile.ZipFile(path) as archive:
                entries = archive.infolist()
                if len(entries) > 10_000 or sum(item.file_size for item in entries) > MAX_EXPANDED_BYTES:
                    raise ValueError("Workbook expanded size exceeds the safe extraction limit")
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                if len(workbook.worksheets) > 200:
                    raise ValueError("Workbook has too many worksheets")
                for sheet in workbook.worksheets:
                    if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                        raise ValueError("Workbook cell dimensions exceed the safe extraction limit")
                    rows = []
                    for row in sheet.iter_rows():
                        total += len(row)
                        if total > MAX_CELLS:
                            raise ValueError("Workbook exceeds the safe cell extraction limit")
                        # Formulas without evaluated cached strings are excluded.
                        rows.append([cell.value if not (isinstance(cell.value, str) and cell.value.startswith("=")) else None for cell in row])
                    sheets.append((sheet.title, rows))
            finally:
                workbook.close()
        else:
            import xlrd

            workbook = xlrd.open_workbook(path, on_demand=True)
            try:
                if workbook.nsheets > 200:
                    raise ValueError("Workbook has too many worksheets")
                for sheet in workbook.sheets():
                    total += sheet.nrows * sheet.ncols
                    if total > MAX_CELLS:
                        raise ValueError("Workbook exceeds the safe cell extraction limit")
                    sheets.append((sheet.name, [sheet.row_values(row) for row in range(sheet.nrows)]))
            finally:
                workbook.release_resources()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Unable to read client brief workbook; upload a valid, unencrypted .xlsx or .xls file") from exc
    return sheets


def parse_brief(path: str, filename: str) -> dict:
    """Extract candidates only from explicit SMS contexts, retaining every source cell."""
    result = {"filename": filename, "sheets": [], "warnings": [], "candidates": []}
    seen = set()

    def add(sheet, row, col, text, campaign="", segment="", language="", scope="SMS", kind="channel_rows"):
        if not _copy(text) or (sheet, row, col) in seen:
            return
        if len(result["candidates"]) >= MAX_CANDIDATES:
            raise ValueError("Workbook contains too many SMS candidates; split the brief into smaller files")
        seen.add((sheet, row, col))
        cell = f"{_column(col + 1)}{row + 1}"
        digest = hashlib.sha256(f"{filename}\0{sheet}\0{cell}\0{campaign}\0{text}".encode()).hexdigest()[:12]
        label = "_".join(filter(None, [campaign or sheet, segment, language, cell]))
        ascii_label = unicodedata.normalize("NFKD", label).encode("ascii", "ignore").decode()
        safe = re.sub(r"[^A-Za-z0-9_]+", "_", ascii_label).strip("_").lower() or "sms"
        warnings = _language_warnings(text, language)
        if _TOKEN.search(text):
            warnings.append("Review every bracketed phrase or placeholder as literal text or a typed variable; no values have been invented.")
        result["candidates"].append({
            "id": f"sms_{digest}", "sheet": sheet, "cell": cell, "text": text,
            "campaign": campaign, "segment": segment or sheet, "language": language,
            "scope": scope, "source_kind": kind,
            "suggested_template_name": f"sms_{safe[:43]}_{digest}",
            "suggested_template_type": "3" if "implicit" in _norm(scope) else "",
            "warnings": warnings, "variables": _TOKEN.findall(text),
        })

    for name, rows in _read_sheets(path, filename):
        before = len(result["candidates"])
        normalized = _norm(name)
        if "planner" in normalized or "calendar" in normalized or _channel(name) == "other":
            result["warnings"].append(f"{name}: skipped non-SMS or planner worksheet.")
        elif rows:
            width = max(map(len, rows), default=0)
            rows = [row + [None] * (width - len(row)) for row in rows]
            campaign_cols = sorted({c for row in rows if not any(_norm(v) == "channel" for v in row)
                                    for c, value in enumerate(row) if _norm(value) == "campaign name"})
            if campaign_cols:
                _campaign_blocks(name, rows, campaign_cols, add)
            else:
                _channel_blocks(name, rows, add)
            if len(result["candidates"]) == before:
                if "regional" in normalized:
                    result["warnings"].append(f"{name}: ambiguous regional source has no explicit SMS channel for its message matrix; skipped rather than assuming SMS approval.")
                else:
                    result["warnings"].append(f"{name}: no explicitly scoped SMS copy found.")
            elif "regional" in normalized and any(_language(row[0]) for row in rows) and not any(_norm(row[0]) in {"lang", "language"} for row in rows):
                result["warnings"].append(f"{name}: unlabelled language-row sections outside an explicit SMS block are skipped as ambiguous/non-SMS.")
        result["sheets"].append({"name": name, "candidate_count": len(result["candidates"]) - before})
    return result


def _campaign_blocks(name, rows, anchors, add):
    """Each label/value block owns its channel, campaign and continuation state."""
    width = len(rows[0])
    for index, anchor in enumerate(anchors):
        end = anchors[index + 1] if index + 1 < len(anchors) else width
        columns = list(range(anchor + 1, end))
        scopes = {c: name if _channel(name) == "sms" else "" for c in columns}
        campaigns = {}
        segment = ""
        active = False
        for ri, row in enumerate(rows):
            label = _norm(row[anchor])
            channel = _channel(row[anchor])
            if channel:
                for c in columns:
                    if label not in {"sms text", "sms utility"} or not scopes[c]:
                        scopes[c] = str(row[anchor]) if channel == "sms" else ""
                active = channel == "sms"
                # Side-by-side variants may override a shared label column's scope.
                for c in columns:
                    if _channel(row[c]):
                        scopes[c] = str(row[c]) if _channel(row[c]) == "sms" else ""
                if label not in {"sms text", "sms utility"}:
                    continue
            if label == "campaign name":
                campaigns = {c: str(row[c]) for c in columns if row[c] is not None}
                active = any(scopes.values())
                continue
            if _metadata(row[anchor]):
                active = False
                continue
            if label and not channel and not re.fullmatch(r"c\d+|[a-z]", label):
                active = False
                if len(label) < 60:
                    segment = str(row[anchor])
                continue
            if active:
                for c in columns:
                    if scopes.get(c):
                        add(name, ri, c, row[c], campaigns.get(c, ""), segment, "", scopes[c], "campaign_block")


def _channel_blocks(name, rows, add):
    width = len(rows[0])
    # Language rows with channel columns (Lang | SMS | WA | PN).
    matrix_header = next((ri for ri, row in enumerate(rows)
                          if _norm(row[0]) in {"lang", "language"} and any(_channel(v) == "sms" for v in row[1:])), None)
    if matrix_header is not None:
        channels = {}
        for ri, row in enumerate(rows):
            if _norm(row[0]) in {"lang", "language"}:
                channels = {c: str(v) for c, v in enumerate(row) if _channel(v) == "sms"}
                continue
            language = _language(row[0])
            if language:
                for c, scope in channels.items():
                    add(name, ri, c, row[c], language=language, scope=scope, kind="language_matrix")
            elif any(v is not None for v in row):
                channels = {}
        return

    # Label columns are discovered independently; content/notes cannot switch a neighbour.
    anchors = sorted({c for row in rows for c, value in enumerate(row)
                      if _norm(value) == "channel" or (_channel(value) and c + 1 < width and any(_copy(r[c + 1]) for r in rows))})
    for index, anchor in enumerate(anchors):
        end = anchors[index + 1] if index + 1 < len(anchors) else width
        content_columns = {anchor + 1: ""} if anchor + 1 < end else {}
        variants = {}
        campaign_column = None
        segment_column = None
        scope = ""
        segment = ""
        for ri, row in enumerate(rows):
            label = _norm(row[anchor])
            # Repeated language/variant headers replace, rather than merge, the old map.
            is_header = label in {"channel", "regional"} and any(row[c] is not None for c in range(anchor + 1, end))
            if is_header:
                content_columns = {}
                variants = {}
                campaign_column = None
                segment_column = None
                for c in range(anchor + 1, end):
                    header = _norm(row[c])
                    language = _language(row[c])
                    if header in {"campaign", "campaign name"}:
                        campaign_column = c
                    if header in {"segment", "usecase", "use case"}:
                        segment_column = c
                    if re.fullmatch(r"[a-z]", header):
                        variants[c] = str(row[c])
                    if language or (header and ("content" in header or header in {"text", "body", "message", "sms text"} or re.fullmatch(r"[a-z]", header))):
                        content_columns[c] = language
                scope = ""
                continue
            channel = _channel(row[anchor])
            if channel:
                scope = str(row[anchor]) if channel == "sms" else ""
            elif label:
                # Unknown row labels, including language labels after PN, end the block.
                scope = ""
                if len(label) < 60 and not _metadata(row[anchor]):
                    segment = str(row[anchor])
                continue
            if scope:
                for c, language in content_columns.items():
                    row_segment = str(row[segment_column]) if segment_column is not None and row[segment_column] is not None else segment
                    row_segment = " / ".join(filter(None, [row_segment, variants.get(c, "")]))
                    campaign = str(row[campaign_column]) if campaign_column is not None and row[campaign_column] is not None else ""
                    add(name, ri, c, row[c], campaign=campaign, segment=row_segment, language=language, scope=scope)
            elif not label and ri == 0:
                headings = [str(row[c]) for c in range(anchor + 1, end)
                            if isinstance(row[c], str) and not _metadata(row[c]) and len(row[c]) < 60]
                if headings:
                    segment = headings[0]


def validate_brief_review(brief: object, text: str) -> list[str]:
    """Validate review evidence without trusting client metadata or rewriting final copy."""
    if not isinstance(brief, dict):
        return ["Brief review must be an object"]
    errors = []
    original = brief.get("original_text")
    if not isinstance(original, str) or not original.strip() or len(original) > MAX_TEXT:
        errors.append("Brief original_text must be nonempty source copy within the workbook cell limit")
        original = ""
    sources = brief.get("sources")
    if not isinstance(sources, list) or not 1 <= len(sources) <= MAX_CANDIDATES:
        errors.append("Brief review requires nonempty source references")
    else:
        for source in sources:
            if not isinstance(source, dict) or any(not isinstance(source.get(key), str) or len(source[key]) > 500
                                                     for key in ("sheet", "cell", "campaign", "segment", "language")):
                errors.append("Brief source references must contain string sheet, cell, campaign, segment and language fields")
                break
            if not source["sheet"].strip() or not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", source["cell"]):
                errors.append("Brief source references require a worksheet and valid cell coordinate")
                break
    warnings = brief.get("warnings")
    if not isinstance(warnings, list) or len(warnings) > MAX_CANDIDATES or any(not isinstance(w, str) or len(w) > MAX_TEXT for w in warnings):
        errors.append("Brief warnings must be a list of strings")
    decisions = brief.get("variables")
    tokens = _TOKEN.findall(original)
    literals = Counter()
    if not isinstance(decisions, list) or len(decisions) != len(tokens):
        errors.append("Brief variable decisions must match every original token in occurrence order")
    else:
        for index, (token, decision) in enumerate(zip(tokens, decisions, strict=True), 1):
            if not isinstance(decision, dict) or decision.get("token") != token:
                errors.append(f"Brief variable decision {index} does not match its original source token")
                continue
            kind = decision.get("kind")
            sample = decision.get("sample")
            if not isinstance(kind, str) or kind not in _KINDS:
                errors.append(f"Resolve brief token {index} as literal text or a typed variable")
            if not isinstance(sample, str) or len(sample) > 2000:
                errors.append(f"Brief token {index} sample must be a string within 2000 characters")
            elif isinstance(kind, str) and kind in _KINDS - {"literal"} and not sample.strip():
                errors.append(f"Provide an explicit sample for brief token {index}")
            if kind == "literal":
                literals[token] += 1
    if brief.get("acknowledged") is not True:
        errors.append("Acknowledge the reviewed source, variable choices and final SMS copy")
    if not isinstance(text, str):
        errors.append("Final brief SMS text must be a string")
    else:
        unresolved = Counter(token for token in _TOKEN.findall(text) if not _TYPED.fullmatch(token))
        if any(count > literals[token] for token, count in unresolved.items()):
            errors.append("Final SMS contains unresolved bracketed or legacy tokens not explicitly reviewed as literal text")
    return errors
